# /*
#  * Copyright 2024 KylinSoft Co., Ltd.
#  *
#  * This program is free software: you can redistribute it and/or modify it under
#  * the terms of the GNU General Public License as published by the Free Software
#  * Foundation, either version 3 of the License, or (at your option) any later
#  * version.
#  *
#  * This program is distributed in the hope that it will be useful, but WITHOUT
#  * ANY WARRANTY; without even the implied warranty of MERCHANTABILITY or FITNESS
#  * FOR A PARTICULAR PURPOSE. See the GNU General Public License for more
#  * details.
#  *
#  * You should have received a copy of the GNU General Public License along with
#  * this program. If not, see <https://www.gnu.org/licenses/>.
#  */

from openpyxl import load_workbook
from .base_parser import BaseParser

class EXCELParser(BaseParser):

    def parser(self, file_path) -> str:
        workbook = load_workbook(filename=file_path)
        text = []
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            for row in sheet.iter_rows(values_only=True):
                for cell_value in row:
                    if cell_value is not None:
                        text.append(str(cell_value))
        result = "".join(text)
        return result

    def is_paragraph_recognition_supported(self,file_type:str) -> bool:
        '''
        todo 待测试是否支持段落识别，目前先返回false
        '''
        return False

if __name__ == '__main__':
    file_path = '/home/kylin/wangyan/test-search/向量数据库.xlsx'
    excel_parser = EXCELParser()
    text = excel_parser.parser(file_path)
    print(text)
