# /*
#  * Copyright 2024 KylinSoft Co., Ltd.
#  *
#  * This program is free software: you can redistribute it and/or modify it under
#  * the terms of the GNU General Public License as published by the Free Software
#  * Foundation, either version 3 of the License, or (at your option) any later
#  * version.
#  *
#  * This program is distributed in the hope that it will be useful, but WITHOUT
#  * ANY WARRANTY; without even the implied warranty of MERCHANTABILITY or FITNESS
#  * FOR A PARTICULAR PURPOSE. See the GNU General Public License for more
#  * details.
#  *
#  * You should have received a copy of the GNU General Public License along with
#  * this program. If not, see <https://www.gnu.org/licenses/>.
#  */

from pptx import Presentation
from .base_parser import BaseParser

class PPTXParser(BaseParser):

    def parser(self, file_path) -> str:
        return self.parser_pptx(file_path)

    def parser_pptx(self, file_path) -> str:
        """
        解析PPTX文件，返回文本内容。
        :param file_path: PPTX文件路径
        :return: 文本内容
        """
        prs = Presentation(file_path)
        content = []

        for slide in prs.slides:
            slide_content = self._parse_slide(slide)
            if slide_content == ['']:
                continue
            else:
                content.extend(slide_content)
        texts = ",".join(content)
        return texts

    def _extract_text_from_shapes(self, shapes, shapes_with_text):
        for shape in shapes:
            if shape.shape_type == 6:  # GroupShape
                self._extract_text_from_shapes(shape.shapes, shapes_with_text)
            elif shape.has_text_frame:
                shapes_with_text.append((shape.top, shape.left, shape.text))

    def _parse_slide(self, slide):
        texts = []
        shapes_with_text = []
        self._extract_text_from_shapes(slide.shapes, shapes_with_text)
        shapes_with_text.sort()
        for _, _, text in shapes_with_text:
            # 去除每个文本框内部的多余空格和换行
            clean_text = " ".join(text.replace('\n', ' ').split())
            if clean_text:
                texts.append(clean_text)

        tables = []
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            for row in shape.table.rows:
                row_data = [" ".join(cell.text.replace('\n', ' ').split()) for cell in row.cells]
                tables.append(", ".join(row_data))

        # 文本框之间用换行拼接，去除多余空格
        result = "\n".join(texts + tables)
        return [result]

    def is_paragraph_recognition_supported(self,file_type:str) -> bool:
        '''
        todo 待测试是否支持段落识别，目前先返回false
        '''
        return False

if __name__ == '__main__':
    parser = PPTXParser()
    print(parser.parser('/home/kylin/prj/prj/v11/scp/kylin-ai-document-service/新能源汽车的发展.pptx'))
